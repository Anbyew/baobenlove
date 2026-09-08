#!/usr/bin/env python3
"""Regenerate the 'Envelope Address' column in testsheet.csv and rebuild testsheet.xlsx.

Usage: python3 build_envelope_addresses.py
Run from anywhere; paths are relative to this script's directory.
"""
import csv
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

DIR = os.path.dirname(os.path.abspath(__file__))
CSV_PATH = os.path.join(DIR, "testsheet.csv")
XLSX_PATH = os.path.join(DIR, "testsheet.xlsx")

COL = "Envelope Address"

COLUMN_WIDTHS = {
    "Name": 22,
    "Name on Envelope": 30,
    "Street Address": 22,
    "Street Address (line 2)": 20,
    "City": 14,
    "State/Region": 12,
    "Zip Code": 10,
    "Country (if not US)": 16,
    "Envelope Address": 32,
}


def build_envelope_address(row):
    name = row["Name on Envelope"].strip()
    name = name.replace(" and ", "\nand ", 1)

    line1 = row["Street Address"].strip()
    line2 = row["Street Address (line 2)"].strip()
    city = row["City"].strip()
    state = row["State/Region"].strip()
    zipc = row["Zip Code"].strip()
    country = row["Country (if not US)"].strip()

    city_line = ", ".join(
        p for p in [city, " ".join(p2 for p2 in [state, zipc] if p2)] if p
    )

    if not line1 and not city_line:
        return name

    lines = [name, "", line1]
    if line2:
        lines.append(line2)
    lines.append(city_line)
    if country:
        lines.append(country)
    return "\n".join(lines)


def update_csv():
    with open(CSV_PATH, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        fieldnames = reader.fieldnames

    if COL not in fieldnames:
        fieldnames = fieldnames + [COL]

    for row in rows:
        row[COL] = build_envelope_address(row)

    with open(CSV_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    return rows


def build_xlsx():
    with open(CSV_PATH, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))

    wb = Workbook()
    ws = wb.active
    ws.title = "Invitations"

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = Alignment(
                wrap_text="\n" in val, vertical="top"
            )

    for c_idx in range(1, len(rows[0]) + 1):
        c = ws.cell(row=1, column=c_idx)
        c.font = c.font.copy(bold=True)

    header = rows[0]
    for c_idx, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = COLUMN_WIDTHS.get(
            name, 18
        )

    for r_idx, row in enumerate(rows[1:], start=2):
        max_lines = max((val.count("\n") + 1 for val in row), default=1)
        ws.row_dimensions[r_idx].height = 15 * max_lines

    ws.freeze_panes = "A2"
    wb.save(XLSX_PATH)


if __name__ == "__main__":
    rows = update_csv()
    build_xlsx()
    print(f"Updated {len(rows)} rows in testsheet.csv and rebuilt testsheet.xlsx")
